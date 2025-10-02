import os
import tempfile
from typing import Dict, Any
import win32com.client as win32
from docx import Document
import mammoth
import io


def extract_doc_content_as_html(file_path: str) -> str:
    """Extract HTML content from .doc file using Windows COM"""
    try:
        word = win32.Dispatch('Word.Application')
        word.Visible = False
        doc = word.Documents.Open(file_path)

        # Save as HTML to temporary file
        temp_html = tempfile.NamedTemporaryFile(suffix='.html', delete=False)
        temp_html.close()

        # Save as Web Page (unfiltered, preserves formatting)
        doc.SaveAs2(temp_html.name, FileFormat=8)  # 8 = wdFormatHTML
        doc.Close()
        word.Quit()

        # Read the HTML content with proper encoding detection
        try:
            with open(temp_html.name, 'r', encoding='utf-8') as f:
                html_content = f.read()
        except UnicodeDecodeError:
            # Try with Windows-1252 encoding
            try:
                with open(temp_html.name, 'r', encoding='windows-1252') as f:
                    html_content = f.read()
            except UnicodeDecodeError:
                # Try with utf-16
                with open(temp_html.name, 'r', encoding='utf-16') as f:
                    html_content = f.read()

        # Clean up temp file
        os.unlink(temp_html.name)

        return html_content
    except Exception as e:
        raise Exception(f"Failed to extract .doc content as HTML: {str(e)}")

def extract_doc_content(file_path: str) -> str:
    """Extract text content from .doc file using Windows COM"""
    try:
        word = win32.Dispatch('Word.Application')
        word.Visible = False
        doc = word.Documents.Open(file_path)
        content = doc.Content.Text
        doc.Close()
        word.Quit()
        return content
    except Exception as e:
        raise Exception(f"Failed to extract .doc content: {str(e)}")


def extract_docx_content(file_path: str) -> str:
    """Extract text content from .docx file"""
    try:
        doc = Document(file_path)
        content = []
        for paragraph in doc.paragraphs:
            content.append(paragraph.text)
        return '\n'.join(content)
    except Exception as e:
        raise Exception(f"Failed to extract .docx content: {str(e)}")


def extract_pdf_content(file_path: str) -> str:
    """Extract text content from PDF file - TODO: Implement later"""
    raise Exception("PDF processing not implemented yet")


def convert_contract_to_html(file_path: str, employee_data: Dict[str, Any]) -> str:
    """
    Convert contract file to HTML with employee data substitution

    Args:
        file_path: Path to the contract file
        employee_data: Dictionary containing employee data for substitution

    Returns:
        HTML string with employee data substituted
    """
    file_extension = os.path.splitext(file_path)[1].lower()

    # Extract content based on file type
    if file_extension == '.doc':
        html_content = extract_doc_content_as_html(file_path)
    elif file_extension == '.docx':
        # Use mammoth for .docx files
        with open(file_path, 'rb') as docx_file:
            result = mammoth.convert_to_html(docx_file)
            html_content = result.value
    elif file_extension == '.pdf':
        content = extract_pdf_content(file_path)
        html_content = content_to_html(content)
    else:
        raise Exception(f"Unsupported file type: {file_extension}")

    # Replace all image src references with base64 encoded sign-stamp image
    import re
    import base64

    # Read and encode image as base64
    try:
        img_path = os.path.join(os.path.dirname(__file__), 'static', 'img', 'sign-stamp.png')
        with open(img_path, 'rb') as img_file:
            img_data = base64.b64encode(img_file.read()).decode('utf-8')
            img_base64 = f'data:image/png;base64,{img_data}'
    except Exception as e:
        print(f"Warning: Could not load sign-stamp.png: {e}")
        img_base64 = '/static/img/sign-stamp.png'  # Fallback to relative path

    html_content = re.sub(r'src="[^"]*\.(gif|jpg|jpeg|png|bmp)"', f'src="{img_base64}"', html_content)

    # Replace placeholders with employee data
    for key, value in employee_data.items():
        placeholder = f"{{{key}}}"
        if placeholder in html_content:
            html_content = html_content.replace(placeholder, str(value) if value else "")

    return html_content


def content_to_html(content: str) -> str:
    """Convert plain text content to HTML with basic formatting"""
    # Split into paragraphs
    paragraphs = content.split('\n')

    html_parts = ['<!DOCTYPE html>', '<html>', '<head>',
                  '<meta charset="UTF-8">',
                  '<title>Contract Document</title>',
                  '<style>',
                  'body { font-family: Arial, sans-serif; margin: 40px; line-height: 1.6; }',
                  'h1 { text-align: center; font-size: 18px; font-weight: bold; margin-bottom: 20px; }',
                  'p { margin-bottom: 10px; text-align: justify; }',
                  '.contract-header { text-align: center; font-weight: bold; margin-bottom: 20px; }',
                  '.article-header { font-weight: bold; margin-top: 20px; margin-bottom: 10px; }',
                  '</style>',
                  '</head>', '<body>']

    for paragraph in paragraphs:
        paragraph = paragraph.strip()
        if not paragraph:
            continue

        # Check if this is a title/header
        if paragraph.startswith('PERJANJIAN KERJA WAKTU TERTENTU'):
            html_parts.append(f'<h1>{paragraph}</h1>')
        elif paragraph.startswith('PASAL'):
            html_parts.append(f'<div class="article-header">{paragraph}</div>')
        elif len(paragraph) < 100 and paragraph.isupper():
            html_parts.append(f'<div class="contract-header">{paragraph}</div>')
        else:
            html_parts.append(f'<p>{paragraph}</p>')

    html_parts.extend(['</body>', '</html>'])

    return '\n'.join(html_parts)


def process_contract_file(file_content: bytes, filename: str, employee_data: Dict[str, Any]) -> str:
    """
    Process uploaded contract file and convert to HTML

    Args:
        file_content: Binary content of the uploaded file
        filename: Original filename
        employee_data: Employee data for placeholder substitution

    Returns:
        HTML string ready for database storage
    """
    # Create temporary file
    with tempfile.NamedTemporaryFile(delete=False, suffix=os.path.splitext(filename)[1]) as temp_file:
        temp_file.write(file_content)
        temp_file_path = temp_file.name

    try:
        # Convert to HTML
        html_content = convert_contract_to_html(temp_file_path, employee_data)
        return html_content
    finally:
        # Clean up temporary file
        os.unlink(temp_file_path)